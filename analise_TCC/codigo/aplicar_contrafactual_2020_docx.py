# -*- coding: utf-8 -*-
"""
Acrescenta 2020 ao exercício contrafactual da seção 4.7.5, à Figura 12 e à Tabela H.1 (14/10/2026).
Números de resultados/contrafactual_anual.xlsx (contrafactual_anual.py, agora com 2020) e, para a
versão com o índice de rigor, de resultados/comparacao_2020_pandemia.xlsx. Tudo conferido com assert.
Backup feito antes: historico/3_versoes_e_backups/backups_do_TCC/_backup_TCC_pre_contrafactual_2020.docx
"""
import copy, io, sys
from pathlib import Path
import numpy as np, pandas as pd
import docx
from docx.oxml.ns import qn
from PIL import Image

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parents[1]
ARQ = RAIZ / "TCC_PEDRO_MINELLA.docx"
RES = RAIZ / "analise_TCC" / "resultados"
FIG = RAIZ / "analise_TCC" / "figuras" / "contrafactual_anual.png"
sys.path.insert(0, str(AQUI))
import lp_base_nova as L


def n(x, casas=2, sinal=False):
    s = (f"%+.{casas}f" if sinal else f"%.{casas}f") % x
    return s.replace(".", ",").replace("-", "−")


# ---------------------------------------------------------------- números
pa = pd.read_excel(RES / "contrafactual_anual.xlsx", sheet_name="por_ano")
tr = pd.read_excel(RES / "contrafactual_anual.xlsx", sheet_name="trajetoria")
tr["Data"] = pd.to_datetime(tr["Data"])


def linha(ano, inicio):
    return pa[(pa.ano == ano) & (pa.grandeza.str.startswith(inicio))].iloc[0]


ob20, cf20, df20 = linha(2020, "observado"), linha(2020, "contrafactual"), linha(2020, "diferenca")
df21, df22 = linha(2021, "diferenca"), linha(2022, "diferenca")
t20 = tr[tr.Data.dt.year == 2020].set_index("Data")
maio = float(t20.loc["2020-05-01", "observado"])
minimo = t20.observado.idxmin()

b, _ = L.carregar("corrigida")
soma_choques = float(b.set_index("Data")["choque_petroleo"].loc["2020-01-01":"2020-12-01"].sum())

ox = pd.read_excel(RES / "comparacao_2020_pandemia.xlsx", sheet_name="contrafactual")
ox = ox[ox.tratamento == "Índice de rigor de Oxford"].set_index("ano")["diferenca"]

# 2021 e 2022 não podem mudar
assert [n(df21.ponto_pp), n(df22.ponto_pp)] == ["1,72", "1,15"]
assert [n(df21.ic5_hac), n(df21.ic95_hac), n(df22.ic5_hac), n(df22.ic95_hac)] == ["0,15", "3,42", "−0,50", "2,88"]
# 2020
assert [n(ob20.ponto_pp), n(cf20.ponto_pp), n(df20.ponto_pp)] == ["−2,27", "0,04", "−2,31"]
assert [n(df20.ic5_hac), n(df20.ic95_hac), n(df20.p_hac, 3)] == ["−5,06", "0,30", "0,143"]
assert [n(ob20.p_hac, 3), n(cf20.p_hac, 3)] == ["0,079", "0,952"]
assert [n(df20.ipca_ano_pontos_log), n(df20.ipca_ano_pct)] == ["4,42", "4,52"]
assert [n(df20.pct_da_inflacao_do_ano, 0), n(ob20.pct_da_inflacao_do_ano, 1),
        n(cf20.pct_da_inflacao_do_ano, 1), n(df20.pct_da_inflacao_do_ano, 1)] == ["−52", "−51,4", "1,0", "−52,4"]
assert minimo == pd.Timestamp("2020-05-01") and n(maio) == "−3,56"
assert float(t20.loc["2020-06-01", "observado"]) > maio          # a recuperação começa em junho
assert n(soma_choques) == "−2,92"
assert [n(ox[2020]), n(ox[2021]), n(ox[2022])] == ["−0,63", "1,10", "0,58"]

# ---------------------------------------------------------------- textos
NOVO_2020 = (
    f"Em 2020 o sinal se inverte. Os choques do ano somam {n(soma_choques)} desvios-padrão, com os dois maiores "
    "choques negativos da amostra em março e abril, e a contribuição sob a paridade chega a "
    f"{n(maio)} pontos em maio, recupera-se em parte a partir de junho e termina o ano em {n(ob20.ponto_pp)}, "
    f"contra {n(cf20.ponto_pp)} sob a regra anterior. A diferença, de {n(df20.ponto_pp)} pontos, corresponde a "
    f"{n(df20.pct_da_inflacao_do_ano, 0)}% da variação do IPCA no ano, de {n(df20.ipca_ano_pontos_log)} pontos "
    f"logarítmicos, ou {n(df20.ipca_ano_pct)}%. Calculada na amostra comum, a contribuição efetivamente ocorrida é "
    f"significativa a 10%, com p de {n(ob20.p_hac, 3)}, mas a diferença não se distingue de zero, com intervalo de "
    f"90% entre {n(df20.ic5_hac)} e {n(df20.ic95_hac)} pontos e p de {n(df20.p_hac, 3)}. O número de 2020 exige "
    "mais cautela que os de 2021 e 2022, por duas razões. Os coeficientes aplicados foram estimados com os dez "
    "meses da pandemia neutralizados pelas indicadoras (seção 4.7.2), de modo que o exercício aplica a 2020 uma "
    "resposta estimada sem ele; com o índice de rigor, que devolve esses meses à estimação, a diferença em 2020 "
    f"cai para {n(ox[2020])} ponto, e as de 2021 e 2022 para {n(ox[2021])} e {n(ox[2022])}. E os coeficientes são "
    "simétricos, ao passo que o Apêndice K.4 documenta, sob a paridade, resposta do diesel bem menor a quedas que "
    "a altas, o que, como ali se registra, leva o exercício a superestimar a contribuição das quedas, que dominam "
    "2020. A leitura que se sustenta é qualitativa: pelo modelo, os choques de 2020 reduziram o IPCA sob a "
    "paridade, e o tamanho dessa redução é incerto."
)

doc = docx.Document(ARQ)
P = doc.paragraphs


def achar(inicio, estilo=None):
    hits = [p for p in P if p.text.startswith(inicio) and (estilo is None or p.style.name == estilo)]
    return hits


def trocar_trecho(p, velho, novo):
    for r in p.runs:
        if velho in r.text:
            r.text = r.text.replace(velho, novo)
            return
    raise ValueError(f"trecho não encontrado em um único run: {velho!r}")


def trocar_tudo(p, novo):
    p.runs[0].text = novo
    for r in p.runs[1:]:
        r.text = ""


# legendas e entradas das listas
cap_fig = achar("Figura 12 – Contribuição dos choques de petróleo ao IPCA sob as duas regras")
assert len(cap_fig) == 2, len(cap_fig)
for p in cap_fig:
    trocar_trecho(p, "por ano civil, 2021 e 2022", "por ano civil, 2020 a 2022")
cap_tab = achar("Tabela H.1 – Contribuição dos choques de petróleo ao IPCA por ano civil")
assert len(cap_tab) == 2, len(cap_tab)
for p in cap_tab:
    trocar_trecho(p, "por ano civil, 2021 e 2022", "por ano civil, 2020 a 2022")

# abertura da seção 4.7.5
[p405] = achar("A magnitude do resultado pode ser dimensionada por um exercício simples.")
trocar_trecho(p405, "período de maior pressão inflacionária da amostra, e acumula-se",
              "período de maior pressão inflacionária da amostra, e, para contraste, os de 2020, ano de queda do "
              "petróleo na pandemia, e acumula-se")

# ressalva de equilíbrio parcial
[p408] = achar("O exercício exige uma ressalva que não pode ser omitida.")
trocar_trecho(p408, " dos choques de petróleo teria sido menor, mantido todo o resto no caminho observado.",
              " dos choques de petróleo teria sido menor em 2021 e 2022, e menos negativa em 2020, mantido todo o "
              "resto no caminho observado.")
trocar_trecho(p408, "que a inflação teria sido 1,72 ponto menor em 2021, nem 1,15 ponto menor em 2022",
              "que a inflação teria sido 2,31 pontos maior em 2020, 1,72 ponto menor em 2021 ou 1,15 ponto menor "
              "em 2022")

# parágrafo novo sobre 2020, logo depois da ressalva e antes da legenda da Figura 12
assert not any(p.text.startswith("Em 2020 o sinal se inverte") for p in P), "parágrafo de 2020 já existe"
novo = copy.deepcopy(p408._p)
p408._p.addnext(novo)
par = docx.text.paragraph.Paragraph(novo, p408._parent)
trocar_tudo(par, NOVO_2020)

# nota da Figura 12
P = doc.paragraphs
[nota_fig] = achar("Nota: contribuição acumulada no ano, reiniciada em janeiro, dos choques ocorridos no próprio ano")
assert "Em 2020" not in nota_fig.text
nota_fig.runs[-1].text = nota_fig.runs[-1].text + (
    " Em 2020, os coeficientes aplicados foram estimados com os meses da pandemia neutralizados pelas indicadoras "
    "(seção 4.7.2).")

# imagem da Figura 12: o parágrafo seguinte à legenda do corpo
i_cap = [i for i, p in enumerate(P) if p.text.startswith("Figura 12 – ") and P[i + 1]._p.xpath(".//a:blip")]
assert len(i_cap) == 1, i_cap
pimg = P[i_cap[0] + 1]
blip = pimg._p.xpath(".//a:blip/@r:embed")[0]
parte = doc.part.related_parts[blip]
novo_png = FIG.read_bytes()
parte._blob = novo_png
w, h = Image.open(io.BytesIO(novo_png)).size
for ext in pimg._p.xpath(".//wp:extent") + pimg._p.xpath(".//a:ext"):
    ext.set("cy", str(int(int(ext.get("cx")) * h / w)))
for pr in pimg._p.xpath(".//wp:docPr"):
    pr.set("descr", "Contribuição dos choques de petróleo ao IPCA por ano civil, 2020 a 2022")

# Tabela H.1: nova coluna de 2020 entre a primeira coluna e a de 2021
tab = [t for t in doc.tables if t.rows[0].cells[1].text == "2021" and t.rows[0].cells[2].text == "2022"
       and t.rows[0].cells[0].text == "Parametrização"]
assert len(tab) == 1, len(tab)
tab = tab[0]
assert tab.rows[3].cells[1].text == "+1,72 p.p. (18,0%)"
vals = ["2020",
        f"{n(ob20.ponto_pp, sinal=True)} p.p. ({n(ob20.pct_da_inflacao_do_ano, 1)}%)",
        f"{n(cf20.ponto_pp, sinal=True)} p.p. ({n(cf20.pct_da_inflacao_do_ano, 1)}%)",
        f"{n(df20.ponto_pp, sinal=True)} p.p. ({n(df20.pct_da_inflacao_do_ano, 1)}%)",
        f"{n(df20.ipca_ano_pontos_log)} ({n(df20.ipca_ano_pct)}%)"]
assert vals[1:4] == ["−2,27 p.p. (−51,4%)", "+0,04 p.p. (1,0%)", "−2,31 p.p. (−52,4%)"], vals
assert vals[4] == "4,42 (4,52%)"

tbl = tab._tbl
grid = tbl.find(qn("w:tblGrid"))
cols = grid.findall(qn("w:gridCol"))
total = sum(int(c.get(qn("w:w"))) for c in cols)
larg = [3420] + [(total - 3420) // 3] * 3
larg[-1] = total - sum(larg[:-1])
novo_col = copy.deepcopy(cols[1])
cols[1].addprevious(novo_col)
for c, wd in zip(grid.findall(qn("w:gridCol")), larg):
    c.set(qn("w:w"), str(wd))
for row, v in zip(tab.rows, vals):
    tcs = row._tr.findall(qn("w:tc"))
    novo_tc = copy.deepcopy(tcs[1])
    tcs[1].addprevious(novo_tc)
    for tc, wd in zip(row._tr.findall(qn("w:tc")), larg):
        tcw = tc.find(qn("w:tcPr")).find(qn("w:tcW"))
        tcw.set(qn("w:w"), str(wd))
    cel = docx.table._Cell(novo_tc, tab)
    trocar_tudo(cel.paragraphs[0], v)

# nota da Tabela H.1 e comentário do Apêndice H
[nota_tab] = achar("Nota: contribuição obtida aplicando-se os coeficientes não acumulados de cada regime")
nota_tab.runs[-1].text = nota_tab.runs[-1].text + " Para 2020, ver a ressalva da seção 4.7.5."
[coment] = achar("As trajetórias mensais constam da Figura 12.")
coment.runs[-1].text = coment.runs[-1].text + (
    f" Em 2020 a contribuição sob a paridade atinge {n(maio)} pontos em maio, depois dos choques negativos de março "
    f"e abril, e termina o ano em {n(ob20.ponto_pp)}.")

doc.save(ARQ)
print("salvo", ARQ)
