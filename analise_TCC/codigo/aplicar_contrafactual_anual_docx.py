# -*- coding: utf-8 -*-
"""
Atualiza a secao 4.7.5, a Figura 12 e a Tabela H.1 do TCC para o contrafactual por ano civil
(07/10/2026). Numeros de resultados/contrafactual_anual.xlsx (contrafactual_anual.py).
Backup feito antes: historico/3_versoes_e_backups/backups_do_TCC/_backup_TCC_pre_contrafactual_anual.docx
"""
import io
from pathlib import Path
import docx
from PIL import Image

RAIZ = Path(__file__).resolve().parents[2]
ARQ = RAIZ / "TCC_PEDRO_MINELLA.docx"
FIG = RAIZ / "analise_TCC" / "figuras" / "contrafactual_anual.png"

d = docx.Document(ARQ)
P = d.paragraphs


def trocar_tudo(p, novo):
    """Texto inteiro no primeiro run (mantem a formatacao dele), demais runs esvaziados."""
    p.runs[0].text = novo
    for r in p.runs[1:]:
        r.text = ""


def trocar_trecho(p, velho, novo):
    for r in p.runs:
        if velho in r.text:
            r.text = r.text.replace(velho, novo)
            return
    raise ValueError(f"trecho nao encontrado em um unico run: {velho!r}")


def conferir(i, inicio):
    assert P[i].text.startswith(inicio), (i, P[i].text[:80])


conferir(404, "A magnitude do resultado pode ser dimensionada")
trocar_tudo(P[404], (
    "A magnitude do resultado pode ser dimensionada por um exercício simples. Tomam-se os choques "
    "efetivamente ocorridos em 2021 e em 2022, período de maior pressão inflacionária da amostra, e "
    "acumula-se a contribuição deles ao IPCA sob duas parametrizações: os coeficientes não acumulados do "
    "regime posterior, que de fato vigoraram, e os do anterior, que teriam vigorado caso a regra não "
    "houvesse mudado. A conta é feita por ano civil e reinicia em janeiro, como a inflação anual. Em 2021, "
    "os choques do ano contribuíram com 1,86 ponto percentual sob a paridade; sob a regra anterior, a "
    "contribuição teria sido de 0,14 ponto, e a diferença de 1,72 ponto corresponde a 18% da variação do "
    "IPCA no ano, de 9,59 pontos logarítmicos, ou 10,06%. Em 2022, as contribuições são de 0,94 e −0,22 "
    "ponto, e a diferença de 1,15 ponto corresponde a 21% da variação de 5,62 pontos logarítmicos, ou "
    "5,79%. Somadas, as diferenças anuais chegam a 2,87 pontos, abaixo dos 3,35 que a acumulação contínua "
    "do biênio produziria, porque o reinício em janeiro deixa de atribuir a 2022 os efeitos dos choques do "
    "fim de 2021. O número exige banda, e ela é larga. Calculada na amostra comum, a diferença tem "
    "intervalo de 90% entre 0,15 e 3,42 pontos em 2021, com p de 0,073, e entre −0,50 e 2,88 pontos em "
    "2022, com p de 0,246. A contribuição efetivamente ocorrida é significativa a 1% em 2021 e não se "
    "distingue de zero em 2022, com p de 0,004 e 0,239; a contrafactual não se distingue de zero em nenhum "
    "dos anos; e a diferença entre as duas mede ordem de grandeza e não deve ser citada como estimativa "
    "pontual."))

conferir(406, "O índice r designa o regime")
trocar_trecho(P[406], "janeiro de 2021 e dezembro de 2022", "janeiro e dezembro de cada ano")
trocar_trecho(P[406], "dentro do biênio", "dentro do próprio ano")

conferir(407, "O exercício exige uma ressalva")
trocar_trecho(P[407], "3,35 pontos menor", "1,72 ponto menor em 2021, nem 1,15 ponto menor em 2022")

for i in (100, 408):
    conferir(i, "Figura 12 – ")
    trocar_tudo(P[i], "Figura 12 – Contribuição dos choques de petróleo ao IPCA sob as duas regras, por ano civil, 2021 e 2022")

conferir(410, "Nota: à esquerda")
trocar_tudo(P[410], (
    "Nota: contribuição acumulada no ano, reiniciada em janeiro, dos choques ocorridos no próprio ano, "
    "calculada com os coeficientes não acumulados de cada regime; a área sombreada é a diferença atribuível "
    "à regra de preços. Acima de cada painel, a diferença ao fim do ano em proporção da variação "
    "logarítmica do IPCA no ano e o seu intervalo de 90%, calculado na amostra comum. O exercício é de "
    "equilíbrio parcial e não constitui projeção da inflação sob regra alternativa."))

for i in (121, 537):
    conferir(i, "Tabela H.1 – ")
    trocar_tudo(P[i], "Tabela H.1 – Contribuição dos choques de petróleo ao IPCA por ano civil, 2021 e 2022")

conferir(538, "Nota: contribuição obtida")
trocar_tudo(P[538], (
    "Nota: contribuição obtida aplicando-se os coeficientes não acumulados de cada regime aos choques "
    "ocorridos em cada ano, com a soma reiniciada em janeiro; entre parênteses, a proporção da variação "
    "logarítmica do IPCA no ano. Exercício de equilíbrio parcial: mede a contribuição dos choques de "
    "petróleo mantido o restante no caminho observado, e não a trajetória da inflação sob regra alternativa."))

conferir(540, "As duas trajetórias mensais")
trocar_tudo(P[540], (
    "As trajetórias mensais constam da Figura 12. Em 2022 a contribuição contrafactual termina levemente "
    "negativa porque, sob os coeficientes do regime anterior, as quedas do preço do petróleo do segundo "
    "semestre pesam mais que as altas do primeiro, ao passo que sob a paridade as altas do primeiro "
    "semestre dominam."))

# Tabela H.1
t = d.tables[4]
assert t.rows[3].cells[1].text == "+3,35 p.p."
novos = [
    [None, "2021", "2022"],
    [None, "+1,86 p.p. (19,4%)", "+0,94 p.p. (16,7%)"],
    [None, "+0,14 p.p. (1,5%)", "−0,22 p.p. (−3,8%)"],
    [None, "+1,72 p.p. (18,0%)", "+1,15 p.p. (20,5%)"],
    ["Variação do IPCA no ano, em pontos logarítmicos (e em %)", "9,59 (10,06%)", "5,62 (5,79%)"],
]
for row, vals in zip(t.rows, novos):
    for cel, v in zip(row.cells, vals):
        if v is not None:
            trocar_tudo(cel.paragraphs[0], v)

# Figura 12: troca a imagem e ajusta a altura mantendo a largura
blip = P[409]._p.xpath(".//a:blip/@r:embed")[0]
parte = d.part.related_parts[blip]
novo = FIG.read_bytes()
parte._blob = novo
w, h = Image.open(io.BytesIO(novo)).size
for ext in P[409]._p.xpath(".//wp:extent") + P[409]._p.xpath(".//a:ext"):
    cx = int(ext.get("cx"))
    ext.set("cy", str(int(cx * h / w)))
for pr in P[409]._p.xpath(".//wp:docPr"):
    if pr.get("descr") is not None:
        pr.set("descr", "Contribuição dos choques de petróleo ao IPCA por ano civil, 2021 e 2022")

d.save(ARQ)
print("salvo", ARQ)
