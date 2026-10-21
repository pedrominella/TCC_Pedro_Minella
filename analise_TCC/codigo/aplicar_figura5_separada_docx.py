# -*- coding: utf-8 -*-
"""
Separa a Figura 5 do TCC em duas figuras e renumera as seguintes (14/10/2026).

  Figura 5  inflação de cada mês após o choque (barras), regime posterior
  Figura 6  resposta acumulada do nível de preços, regime posterior (a da Tabela 2)
  Figuras 6 a 13 passam a 7 a 14: legendas, lista de figuras e citações no texto.

Imagens de fig_distribuicao_temporal.py. A renumeração confere, parágrafo a parágrafo, que o
texto final é o esperado, o que pega citação partida entre trechos de formatação.
Backup feito antes: historico/3_versoes_e_backups/backups_do_TCC/_backup_TCC_pre_figura5_separada.docx
"""
import copy, io, re
from pathlib import Path
import docx
from docx.oxml.ns import qn
from PIL import Image

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parents[1]
ARQ = RAIZ / "TCC_PEDRO_MINELLA.docx"
FIG5 = RAIZ / "analise_TCC" / "figuras" / "distribuicao_temporal_mensal.png"
FIG6 = RAIZ / "analise_TCC" / "figuras" / "distribuicao_temporal_acumulada.png"

CAP5 = "Figura 5 – Distribuição temporal do repasse: inflação de cada mês após o choque"
CAP6 = "Figura 6 – Distribuição temporal do repasse: resposta acumulada do nível de preços"

doc = docx.Document(ARQ)


def todos_paragrafos(d):
    for p in d.paragraphs:
        yield p
    for t in d.tables:
        for row in t.rows:
            for cel in row.cells:
                for p in cel.paragraphs:
                    yield p


assert not any(p.text.startswith("Figura 14 –") for p in doc.paragraphs), "a renumeração já foi feita"

# ---------------------------------------------------------------- 1) renumerar 6..13 -> 7..14
PADRAO = re.compile(r"\b(Figuras?)(\s+)(\d+(?:(?:,\s*|\s+e\s+)\d+)*)")


def renum(m):
    nums = re.sub(r"\d+", lambda k: str(int(k.group()) + 1) if 6 <= int(k.group()) <= 13 else k.group(), m.group(3))
    return m.group(1) + m.group(2) + nums


mudados = 0
vistos = set()
for p in todos_paragrafos(doc):
    if id(p._p) in vistos:
        continue
    vistos.add(id(p._p))
    esperado = PADRAO.sub(renum, p.text)
    if esperado == p.text:
        continue
    for r in p.runs:
        r.text = PADRAO.sub(renum, r.text)
    assert p.text == esperado, ("citação partida entre trechos", p.text[:120])
    mudados += 1
print("parágrafos renumerados:", mudados)

P = doc.paragraphs


def unico(cond, rot):
    achados = [i for i, p in enumerate(P) if cond(p)]
    assert len(achados) == 1, (rot, achados)
    return achados[0]


# ---------------------------------------------------------------- 2) Figura 5: legenda, imagem, nota
i_lista5 = [i for i, p in enumerate(P) if p.text == "Figura 5 – Distribuição temporal do repasse"]
assert len(i_lista5) == 2, i_lista5            # entrada da lista e legenda do corpo
i_cap5 = [i for i in i_lista5 if P[i + 1]._p.xpath(".//a:blip")]
assert len(i_cap5) == 1
i_cap5 = i_cap5[0]
i_ent5 = [i for i in i_lista5 if i != i_cap5][0]
for i in (i_ent5, i_cap5):
    P[i].runs[0].text = CAP5
    for r in P[i].runs[1:]:
        r.text = ""

pimg5 = P[i_cap5 + 1]
nota5 = P[i_cap5 + 2]
fonte5 = P[i_cap5 + 3]
assert nota5.text.startswith("Nota: regime posterior a março de 2017. As barras representam")
assert fonte5.text == "Fonte: Elaboração própria."


def trocar_imagem(pimg, arq, descr):
    rid = pimg._p.xpath(".//a:blip/@r:embed")[0]
    parte = doc.part.related_parts[rid]
    novo = arq.read_bytes()
    parte._blob = novo
    w, h = Image.open(io.BytesIO(novo)).size
    for ext in pimg._p.xpath(".//wp:extent") + pimg._p.xpath(".//a:ext"):
        ext.set("cy", str(int(int(ext.get("cx")) * h / w)))
    for pr in pimg._p.xpath(".//wp:docPr"):
        pr.set("descr", descr)


trocar_imagem(pimg5, FIG5, "Inflação de cada mês após o choque de petróleo, regime posterior a março de 2017")
runs = nota5.runs
assert len(runs) == 3 and runs[1].text == "t+h", [r.text for r in runs]
runs[2].text = (", estimada separadamente para cada mês, com intervalo de 90%; em verde, contribuições positivas, "
                "e em vermelho, negativas. A resposta acumulada está na Figura 6.")

# ---------------------------------------------------------------- 3) Figura 6 nova, depois da fonte da Figura 5
cap6 = copy.deepcopy(P[i_cap5]._p)
img6 = copy.deepcopy(pimg5._p)
nota6 = copy.deepcopy(nota5._p)
fonte6 = copy.deepcopy(fonte5._p)
ancora = fonte5._p
for el in (cap6, img6, nota6, fonte6):
    ancora.addnext(el)
    ancora = el

pc6 = docx.text.paragraph.Paragraph(cap6, fonte5._parent)
pc6.runs[0].text = CAP6
for r in pc6.runs[1:]:
    r.text = ""

rid6, _ = doc.part.get_or_add_image(str(FIG6))
for blip in img6.xpath(".//a:blip"):
    blip.set(qn("r:embed"), rid6)
w6, h6 = Image.open(FIG6).size
for ext in img6.xpath(".//wp:extent") + img6.xpath(".//a:ext"):
    ext.set("cy", str(int(int(ext.get("cx")) * h6 / w6)))
ids = [int(x) for x in doc.element.body.xpath(".//wp:docPr/@id")]
for pr in img6.xpath(".//wp:docPr"):
    pr.set("id", str(max(ids) + 1))
    pr.set("name", f"Imagem {max(ids) + 1}")
    pr.set("descr", "Resposta acumulada do nível de preços ao choque de petróleo, regime posterior a março de 2017")

pn6 = docx.text.paragraph.Paragraph(nota6, fonte5._parent)
r6 = pn6.runs
r6[0].text = "Nota: regime posterior a março de 2017. Resposta acumulada do nível de preços "
r6[1].text = "h"
r6[2].text = " meses após o choque, a mesma da Tabela 2, com faixa de 90%."

# entrada na lista de figuras
ent6 = copy.deepcopy(P[i_ent5]._p)
P[i_ent5]._p.addnext(ent6)
pe6 = docx.text.paragraph.Paragraph(ent6, P[i_ent5]._parent)
pe6.runs[0].text = CAP6
for r in pe6.runs[1:]:
    r.text = ""

# ---------------------------------------------------------------- 4) citação no texto
P = doc.paragraphs
i_txt = unico(lambda p: p.text.endswith("A Figura 5 exibe o padrão."), "citação da Figura 5")
for r in P[i_txt].runs:
    if "A Figura 5 exibe o padrão." in r.text:
        r.text = r.text.replace("A Figura 5 exibe o padrão.",
                                "A Figura 5 mostra a inflação de cada mês, e a Figura 6, a resposta acumulada "
                                "correspondente.")
        break
else:
    raise ValueError("frase da Figura 5 partida entre trechos")

doc.save(ARQ)
print("salvo", ARQ)
