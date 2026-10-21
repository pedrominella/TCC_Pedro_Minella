# -*- coding: utf-8 -*-
"""
Recupera a introdução do arquivo mestre no arquivo de revisão, como alteração controlada (14/10/2026).

Os seis parágrafos atuais da Introdução (versão enxugada) ficam marcados como excluídos e os onze do
mestre entram marcados como inseridos, autor "Claude". O texto do mestre é mantido, com estes ajustes
de coerência com decisões posteriores a ele:
  - resultado principal: horizontes explicitados (impacto, um e três meses) e o intervalo de 90% do
    coeficiente anterior no lugar de "o zero está bem estimado" (onda2_8, Tabela 2); a frase "sob os dois
    tratamentos" foi conferida em 14/10 (indicadoras e Oxford: 18 de 18 posteriores a 5%, nenhum anterior
    a 10%)
  - dicotomia mercado/administrativa: "tende a ser característica mais estável" (resposta ao comentário)
  - política monetária: "dificilmente", e não "a candidata que sobra", coerente com a endogeneidade
  - método: frase da endogeneidade assumida e instrumentos fracos
  - proposições: data condicionada ao tratamento principal da pandemia; canal como parcela associada;
    segunda volta pontual, não significativa em três meses (p = 0,124), sem confirmação independente
  - implicação: "depende da regra", em vez de "atributo da regra"
As âncoras dos comentários dos parágrafos antigos passam para os parágrafos novos correspondentes.

Uso: python aplicar_introducao_mestre_docx.py [entrada saida]
Backup: historico/3_versoes_e_backups/backups_do_TCC/_backup_revisao_pre_introducao.docx
"""
import copy, random, sys
from pathlib import Path
import pandas as pd
import docx
from docx.oxml.ns import qn

AQUI = Path(__file__).resolve().parent
sys.path.insert(0, str(AQUI))
from revisao_track import Revisor

RAIZ = AQUI.parents[1]
RES = AQUI.parent / "resultados"
MESTRE = RAIZ / "TCC_PEDRO_MINELLA.docx"
ENTRADA = Path(sys.argv[1]) if len(sys.argv) > 2 else RAIZ / "TCC_PEDRO_MINELLA_revisao_com_alteracoes.docx"
SAIDA = Path(sys.argv[2]) if len(sys.argv) > 2 else ENTRADA
W14_PARAID = "{http://schemas.microsoft.com/office/word/2010/wordml}paraId"


def br(x, casas=3):
    return f"{x:.{casas}f}".replace(".", ",").replace("-", "−")


def intro(d):
    P = d.paragraphs
    i0 = [i for i, p in enumerate(P) if p.style.name == "Heading 1" and p.text.strip().startswith("1 INTRODU")][0]
    i1 = [i for i, p in enumerate(P) if i > i0 and p.style.name == "Heading 1"][0]
    return P[i0], P[i0 + 1:i1]


# ------------------------------------------------------------------ texto do mestre, conferido
_, mestre = intro(docx.Document(MESTRE))
M = [p.text for p in mestre]
assert len(M) == 11 and M[0].startswith("Conflitos armados no Oriente Médio") and M[10].startswith("Além desta introdução")

t2 = pd.read_excel(RES / "onda2_8_teulings_zubanov.xlsx", sheet_name="respostas")
pre0 = t2[(t2.modo == "regime") & (t2.versao == "sem leads") & (t2.alvo == "choque_pre")
          & (t2.variavel == "ipca_geral") & (t2.h == 0)].iloc[0]
assert round(pre0.erro_padrao, 3) == 0.024 and pre0.pontual_superior < 0.176
ic = f"{br(pre0.pontual_inferior)} a {br(pre0.pontual_superior)}"


def troca(texto, velho, novo):
    assert texto.count(velho) == 1, velho[:60]
    return texto.replace(velho, novo)


N = list(M)
N[1] = troca(N[1], "nas dezoito estimativas do regime anterior nenhuma",
             "nas dezoito estimativas do regime anterior nos horizontes de impacto, um e três meses, nenhuma")
N[1] = troca(N[1], "No índice cheio, o coeficiente anterior tampouco é impreciso, com erro-padrão de 0,024 no impacto, "
                   "de modo que o zero está bem estimado, e o repasse passou de indetectável a sistemático.",
             f"No índice cheio, o coeficiente anterior tampouco é impreciso: com erro-padrão de 0,024 no impacto, "
             f"seu intervalo de 90% vai de {ic} e fica bem abaixo da resposta posterior, de modo que o repasse "
             f"passou de indetectável a sistemático.")
N[2] = troca(N[2], "o repasse é característica razoavelmente permanente da economia",
             "o repasse tende a ser característica mais estável da economia")
N[3] = troca(N[3], "A explicação não está, portanto, na política monetária, e a candidata que sobra é a regra de "
                   "formação de preços.",
             "A explicação dificilmente está na política monetária, e a regra de formação de preços é a candidata "
             "natural, ainda que o desenho do trabalho não permita excluir outras (seção 3.4).")
N[7] = troca(N[7], "o que o torna independente da estimação.", "sem depender da estimação da data.")
N[7] = troca(N[7], "A amostra é mensal, de janeiro de 2003 a setembro de 2025.",
             "A amostra é mensal, de janeiro de 2003 a setembro de 2025. Como os instrumentos externos testados "
             "se mostraram fracos, o choque é tratado como endógeno, e o contraste entre regimes é lido como "
             "efeito da regra sob a hipótese de vieses comparáveis nos dois períodos.")
N[8] = troca(N[8], "data que os dados elegem e a cronologia institucional corrobora.",
             "data que os dados elegem com o tratamento principal da pandemia e que a cronologia institucional "
             "corrobora.")
N[8] = troca(N[8], "Transita pelos preços de combustíveis, canal que carrega 0,003 ponto percentual em três meses no "
                   "regime anterior e 0,249 no posterior, diferença com p inferior a 0,001.",
             "Está associado aos preços de combustíveis, cuja parcela é de 0,003 ponto percentual em três meses no "
             "regime anterior e de 0,249 no posterior, diferença com p inferior a 0,001.")
N[8] = troca(N[8], "E torna-se majoritariamente de segunda volta a partir do terceiro mês, magnitude para a qual duas "
                   "construções distintas convergem, ainda que nenhuma a estabeleça isoladamente.",
             "E, na estimativa pontual, torna-se majoritariamente de segunda volta a partir do terceiro mês, ainda "
             "que a segunda volta não seja significativa em três meses (p = 0,124) e que as duas construções "
             "usadas não sejam confirmações independentes.")
N[9] = troca(N[9], "não é constante da economia, e sim atributo da regra de formação de preços vigente.",
             "não é constante da economia: depende da regra de formação de preços vigente.")
assert not any(c in "".join(N) for c in "—–")

# ------------------------------------------------------------------ revisão
doc = docx.Document(ENTRADA)
R = Revisor(doc)
corpo = doc.element.body
titulo, antigos = intro(doc)
antigos = [p._p for p in antigos]
assert len(antigos) == 6 and docx.text.paragraph.Paragraph(antigos[0], None).text.startswith("A recorrência de choques")
for p in antigos:
    assert p.find(".//" + qn("w:ins")) is None and p.find(".//" + qn("w:del")) is None

usados = {e.get(W14_PARAID) for e in corpo.iter() if e.get(W14_PARAID)}
rnd = random.Random(20261014)


def paraid():
    while True:
        v = "%08X" % rnd.randrange(1, 0x7FFFFFFF)
        if v not in usados:
            usados.add(v)
            return v


def marca_paragrafo(p, tipo):
    ppr = p.find(qn("w:pPr"))
    rpr = ppr.find(qn("w:rPr"))
    if rpr is None:
        rpr = ppr.makeelement(qn("w:rPr"), {}); ppr.append(rpr)
    rpr.insert(0, R._marca(tipo))


modelo_p = antigos[1]                      # parágrafo comum do corpo: recuo, justificado, idioma
modelo_r = modelo_p.find(qn("w:r"))
novos = []
for texto in N:
    p = corpo.makeelement(qn("w:p"), {})
    p.set(W14_PARAID, paraid())
    ppr = copy.deepcopy(modelo_p.find(qn("w:pPr")))
    for velho in ppr.findall(qn("w:rPr")):
        ppr.remove(velho)
    p.append(ppr)
    ins = R._marca("w:ins"); ins.append(R._run_novo(modelo_r, texto)); p.append(ins)
    marca_paragrafo(p, "w:ins")
    novos.append(p)
ancora = titulo._p
for p in novos:
    ancora.addnext(p); ancora = p

# âncoras de comentário: parágrafo antigo -> parágrafo novo correspondente
mapa = {0: 0, 1: 1, 2: 2, 3: 5, 4: 7, 5: 10}
for k, p in enumerate(antigos):
    alvo = novos[mapa[k]]
    inicios = [e for e in p if e.tag == qn("w:commentRangeStart")]
    fins = [e for e in p if e.tag == qn("w:commentRangeEnd")
            or (e.tag == qn("w:r") and e.find(qn("w:commentReference")) is not None)]
    texto = "".join(t.text or "" for t in p.iter(qn("w:t")))
    R.substituir(p, texto, "")
    marca_paragrafo(p, "w:del")
    ppr = alvo.find(qn("w:pPr"))
    for e in reversed(inicios):
        ppr.addnext(e)
    for e in fins:
        alvo.append(e)

doc.save(SAIDA)
print("salvo", SAIDA, "| intervalo do coeficiente anterior:", ic)
