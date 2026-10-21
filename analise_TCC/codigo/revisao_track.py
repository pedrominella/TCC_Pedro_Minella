# -*- coding: utf-8 -*-
"""
Alterações controladas (w:ins/w:del) em docx, para o arquivo de revisão do TCC (14/10/2026).

Mesmas rotinas de aplicar_endogeneidade_revisao_docx.py, em módulo, para as próximas rodadas.
Só mexe em runs fora de revisões existentes e fora de fórmulas; o texto de outras revisões fica
intacto, de modo que rejeitar as marcações do autor AUTOR devolve o parágrafo ao estado anterior.
"""
import copy
from datetime import datetime, timezone
from docx.oxml.ns import qn

AUTOR = "Claude"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
OMATH = "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath"


class Revisor:
    def __init__(self, doc, autor=AUTOR):
        self.doc, self.autor = doc, autor
        self.data = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
        ids = [int(e.get(qn("w:id"))) for e in doc.element.body.iter()
               if e.get(qn("w:id")) is not None and e.get(qn("w:id")).lstrip("-").isdigit()]
        self._id = max(ids + [0]) + 1000

    # ----------------------------------------------------------- utilitários
    def _marca(self, tag):
        self._id += 1
        return self.doc.element.body.makeelement(
            qn(tag), {qn("w:id"): str(self._id), qn("w:author"): self.autor, qn("w:date"): self.data})

    @staticmethod
    def _runs(p):
        saida = []
        for r in p.iter(qn("w:r")):
            if {a.tag for a in r.iterancestors()} & {qn("w:ins"), qn("w:del"), OMATH}:
                continue
            if r.find(qn("w:t")) is not None:
                saida.append(r)
        return saida

    @staticmethod
    def _texto(r):
        return "".join(t.text or "" for t in r.findall(qn("w:t")))

    def _partir(self, r, k):
        s = self._texto(r)
        if k <= 0 or k >= len(s):
            return
        assert len(r.findall(qn("w:t"))) == 1, "run com vários w:t"
        novo = copy.deepcopy(r)
        for filho in list(novo):                 # a parte direita leva só texto: tabs e quebras ficam à esquerda
            if filho.tag not in (qn("w:rPr"), qn("w:t")):
                novo.remove(filho)
        r.find(qn("w:t")).text, novo.find(qn("w:t")).text = s[:k], s[k:]
        for t in (r.find(qn("w:t")), novo.find(qn("w:t"))):
            t.set(XML_SPACE, "preserve")
        r.addnext(novo)

    def _isolar(self, p, trecho):
        texto = "".join(self._texto(r) for r in self._runs(p))
        assert texto.count(trecho) == 1, ("trecho ausente ou repetido", trecho[:70], texto.count(trecho))
        ini = texto.index(trecho); fim = ini + len(trecho)
        for pos in (ini, fim):
            acc = 0
            for r in self._runs(p):
                n = len(self._texto(r))
                if acc < pos < acc + n:
                    self._partir(r, pos - acc)
                    break
                acc += n
        dentro, acc = [], 0
        for r in self._runs(p):
            n = len(self._texto(r))
            if acc >= ini and acc + n <= fim and n:
                dentro.append(r)
            acc += n
        for r in dentro:                          # tabs e quebras não entram na marcação
            self._separar_nao_texto(r)
        return dentro

    @staticmethod
    def _separar_nao_texto(r):
        """Tira de r os filhos que não são texto (tab, quebra) e os põe em runs vizinhos, na mesma ordem."""
        filhos = [f for f in r if f.tag != qn("w:rPr")]
        if all(f.tag == qn("w:t") for f in filhos):
            return
        rpr = r.find(qn("w:rPr"))
        antes, depois, viu_t = [], [], False
        for f in filhos:
            if f.tag == qn("w:t"):
                viu_t = True
            else:
                (depois if viu_t else antes).append(f)
        for grupo, lado in ((antes, "antes"), (depois, "depois")):
            if not grupo:
                continue
            novo = r.makeelement(qn("w:r"), {})
            if rpr is not None:
                novo.append(copy.deepcopy(rpr))
            for f in grupo:
                r.remove(f)
                novo.append(f)
            (r.addprevious if lado == "antes" else r.addnext)(novo)

    def substituir_por(self, p, antigo, elementos):
        """Como substituir, mas o conteúdo novo são elementos prontos (equações, runs)."""
        alvo = self._isolar(p, antigo)
        d = self._marca("w:del")
        alvo[0].addprevious(d)
        for r in alvo:
            d.append(r)
            for t in r.findall(qn("w:t")):
                t.tag = qn("w:delText")
                t.set(XML_SPACE, "preserve")
        i = self._marca("w:ins")
        for el in elementos:
            i.append(el)
        d.addnext(i)
        return i

    def apagar_elemento(self, el):
        """Marca um elemento inteiro (por exemplo uma equação antiga) como excluído."""
        d = self._marca("w:del")
        el.addprevious(d)
        d.append(el)
        return d

    def _run_novo(self, modelo, texto):
        r = copy.deepcopy(modelo)
        for filho in list(r):
            if filho.tag != qn("w:rPr"):
                r.remove(filho)
        t = r.makeelement(qn("w:t"), {})
        t.text = texto
        t.set(XML_SPACE, "preserve")
        r.append(t)
        return r

    # ----------------------------------------------------------- operações
    def substituir(self, p, antigo, novo):
        alvo = self._isolar(p, antigo)
        d = self._marca("w:del")
        alvo[0].addprevious(d)
        for r in alvo:
            d.append(r)
            for t in r.findall(qn("w:t")):
                t.tag = qn("w:delText")
                t.set(XML_SPACE, "preserve")
        if novo:
            i = self._marca("w:ins")
            i.append(self._run_novo(alvo[0], novo))
            d.addnext(i)

    def inserir_depois(self, p, ancora, novo):
        alvo = self._isolar(p, ancora)
        i = self._marca("w:ins")
        i.append(self._run_novo(alvo[-1], novo))
        alvo[-1].addnext(i)

    # ----------------------------------------------------------- localização
    def paragrafo(self, inicio, aceita_revisao=False):
        achados = [p for p in self.doc.paragraphs if p.text.startswith(inicio)]
        assert len(achados) == 1, (inicio[:60], len(achados))
        p = achados[0]._p
        if not aceita_revisao:
            assert p.find(".//" + qn("w:ins")) is None and p.find(".//" + qn("w:del")) is None, \
                ("parágrafo já tem revisão", inicio[:60])
        return p

    def paragrafos_com(self, trecho):
        return [p._p for p in self.doc.paragraphs if trecho in p.text]

    def tabela(self, primeira_celula, contem):
        achadas = [t for t in self.doc.tables
                   if t.rows[0].cells[0].text.strip() == primeira_celula
                   and any(contem in c.text for r in t.rows for c in r.cells)]
        assert len(achadas) == 1, (primeira_celula, contem, len(achadas))
        return achadas[0]

    @staticmethod
    def celula(tabela, i, j):
        ps = tabela.rows[i].cells[j].paragraphs
        assert len(ps) == 1, ("célula com vários parágrafos", i, j)
        return ps[0]._p
