# -*- coding: utf-8 -*-
"""
Capa e pré-textuais (14/10/2026): logo do Ibmec no topo da capa, nome completo do autor
(PEDRO FRANCK MINELLA) na capa, folha de rosto e folha de aprovação, e banca com o orientador
e um examinador (sai o terceiro). Logo: imagem enviada pelo autor, recortada nas bordas brancas
e guardada em documentos_complementares/logo_ibmec_capa.png.
Backup feito antes: historico/3_versoes_e_backups/backups_do_TCC/_backup_TCC_pre_capa_logo.docx
"""
import sys
from pathlib import Path
import numpy as np
from PIL import Image
import docx
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

RAIZ = Path(__file__).resolve().parents[2]
ARQ = RAIZ / "TCC_PEDRO_MINELLA.docx"
ORIG = Path(sys.argv[1])
LOGO = RAIZ / "documentos_complementares" / "logo_ibmec_capa.png"

im = Image.open(ORIG).convert("RGB")
a = np.asarray(im.convert("L"))
ys, xs = np.where(a < 240)
m = 12
im.crop((max(xs.min() - m, 0), max(ys.min() - m, 0), min(xs.max() + m, im.width), min(ys.max() + m, im.height))).save(LOGO)

doc = docx.Document(ARQ)
P = doc.paragraphs
assert P[0].text.strip() == "IBMEC" and P[1].text.strip() == "CURSO DE CIÊNCIAS ECONÔMICAS"
assert [P[i].text for i in (8, 29, 55)] == ["PEDRO MINELLA"] * 3
assert [P[i].text for i in (76, 77, 78)] == ["_" * 52, "[NOME DO EXAMINADOR]", "IBMEC (Examinador)"]
assert P[79].text == "RESUMO" and "pageBreakBefore" in P[79]._p.xml

# banca: sai o terceiro examinador
for i in (78, 77, 76):
    P[i]._p.getparent().remove(P[i]._p)
# nome completo
for i in (8, 29, 55):
    P[i].runs[0].text = "PEDRO FRANCK MINELLA"
    for r in P[i].runs[1:]:
        r.text = ""
# capa: duas linhas em branco a menos, para compensar a altura da logo
vazios = [p for p in P[2:8] if not p.text.strip() and not p._p.xpath(".//a:blip")]
assert len(vazios) >= 2
for p in vazios[:2]:
    p._p.getparent().remove(p._p)
# logo no topo da capa
novo = P[0].insert_paragraph_before()
novo.alignment = WD_ALIGN_PARAGRAPH.CENTER
novo.paragraph_format.space_after = Cm(0.4)
novo.add_run().add_picture(str(LOGO), width=Cm(4.5))
doc.save(ARQ)
print("salvo", ARQ)
