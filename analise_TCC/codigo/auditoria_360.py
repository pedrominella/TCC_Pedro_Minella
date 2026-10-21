# -*- coding: utf-8 -*-
"""
AUDITORIA 360 DO TCC

Motivo de existir: em 05/09/2026 descobriu-se que a Tabela 5 e as Figuras 8 e 9
usavam AMOSTRA CHEIA enquanto o capitulo inteiro e POR REGIME. A auditoria
anterior comparava datas de arquivo e nao pegou, porque as figuras tinham sido
regeradas corretamente -- a partir da planilha errada. Data prova frescor, nao
prova coerencia.

Este script verifica COERENCIA, em seis camadas independentes:

  1 RASTREABILIDADE   todo numero do .docx casa com alguma celula de resultado?
  2 ESPECIFICACAO     cada script estima com qual amostra e qual especificacao?
  3 IDENTIDADES       as relacoes contabeis fecham numericamente?
  4 CRUZADA           a mesma grandeza tem o mesmo valor em todo lugar?
  5 CADEIA            dado -> script -> resultado -> figura -> documento em ordem?
  6 REMISSOES         toda remissao aponta para algo que existe? listas batem?

Sai com codigo 1 se houver ACHADO CRITICO. Avisos nao derrubam.

Uso:  python auditoria_360.py            relatorio completo
      python auditoria_360.py 2          so a camada 2
"""
import os, re, sys, time, warnings
from pathlib import Path
import numpy as np, pandas as pd

warnings.filterwarnings("ignore")

RAIZ = Path(r"C:/Users/pedro/OneDrive/Documentos/TCC_python")
ART = RAIZ / "analise_TCC"
COD, RES, FIG = ART / "codigo", ART / "resultados", ART / "figuras"
DOCX = RAIZ / "TCC_PEDRO_MINELLA.docx"
GEN = Path(r"C:/Users/pedro/AppData/Local/Temp/claude/"
           r"c--Users-pedro-OneDrive-Documentos-TCC-python/"
           r"b948f4bc-20ef-496a-b291-1171ce8ec6b1/scratchpad")

CRIT, AVISO = [], []


def crit(camada, msg):
    CRIT.append((camada, msg)); print("   [CRITICO] %s" % msg)


def aviso(camada, msg):
    AVISO.append((camada, msg)); print("   [aviso]   %s" % msg)


def ok(msg):
    print("   [ok]      %s" % msg)


def cab(n, t):
    print("\n" + "=" * 78); print("CAMADA %s  %s" % (n, t)); print("=" * 78)


# ══════════════════════════════════════════════════ utilidades
def ler_docx():
    from docx import Document
    d = Document(str(DOCX))
    paras = [p.text for p in d.paragraphs]
    cels = [[c.text.strip() for c in r.cells] for t in d.tables for r in t.rows]
    return d, paras, cels


_CACHE = {}


IGNORAR = ("_superseded", "_antes_rerun", "_backup")


def relevantes(pasta, padrao):
    """Arquivos de resultado vivos: exclui pastas de arquivo morto."""
    return [f for f in pasta.rglob(padrao)
            if not any(x in f.parts for x in IGNORAR)]


def todas_celulas():
    """Todo valor numerico de todas as planilhas de resultado, com origem."""
    if _CACHE:
        return _CACHE["vals"]
    vals = {}
    for f in sorted(relevantes(RES, "*.xlsx")):
        try:
            x = pd.ExcelFile(f)
        except Exception:
            continue
        for ab in x.sheet_names:
            try:
                d = x.parse(ab)
            except Exception:
                continue
            for c in d.columns:
                s = pd.to_numeric(d[c], errors="coerce").dropna()
                for v in s.values:
                    vals.setdefault(round(float(v), 3), set()).add("%s:%s" % (f.name, ab))
    _CACHE["vals"] = vals
    return vals


def numeros_pt(txt):
    """Numeros no formato brasileiro com 2 a 4 decimais."""
    return [(m.group(0), float(m.group(0).replace(".", "").replace(",", ".")))
            for m in re.finditer(r"-?\d{1,3}(?:\.\d{3})*,\d{2,4}\b", txt)]


# ══════════════════════════════════════════════════ 1 RASTREABILIDADE
def camada1(paras, cels):
    cab(1, "RASTREABILIDADE: todo numero do documento vem de alguma planilha?")
    vals = todas_celulas()
    print("   base de comparacao: %d valores distintos nas planilhas de resultado"
          % len(vals))
    orfaos, casados = [], 0
    fonte = [(" | ".join(r), "tabela") for r in cels] + [(p, "texto") for p in paras]
    for txt, onde in fonte:
        for bruto, v in numeros_pt(txt):
            achou = any(round(v + d, 3) in vals for d in (0, 0.001, -0.001))
            if achou:
                casados += 1
            else:
                orfaos.append((bruto, onde, txt[:90]))
    print("   numeros casados: %d | orfaos: %d" % (casados, len(orfaos)))
    if orfaos:
        aviso(1, "%d numeros nao casam com celula alguma (podem ser derivados, "
                 "percentuais ou digitados)" % len(orfaos))
        for b, onde, ctx in orfaos[:12]:
            print("             %-9s (%s) ...%s..." % (b, onde, ctx.strip()[:70]))
        if len(orfaos) > 12:
            print("             ... e mais %d" % (len(orfaos) - 12))
    else:
        ok("todo numero do documento tem origem rastreada")


# ══════════════════════════════════════════════════ 2 ESPECIFICACAO
# Padroes que indicam extracao de AMOSTRA CHEIA de um arquivo de resultado.
# Nao sao heuristicas de nome de variavel: sao as formas concretas pelas quais
# um script pega a coluna de amostra cheia. Foi assim que o bug de 05/09 entrou.
CHEIA = [
    (r'modo\s*==\s*["\']cheia["\']', 'filtra modo == "cheia"'),
    (r'alvo\s*==\s*["\']choque_petroleo["\']', 'filtra alvo == "choque_petroleo"'),
    (r'ajustar\([^)]*["\']choque_petroleo["\']\)', 'ajustar() no choque unico'),
    (r'onda2_7_primeira_segunda_volta\.xlsx', 'le a decomposicao de amostra cheia'),
    (r'sheet_name\s*=\s*["\']amostra_cheia["\']', 'le a aba amostra_cheia'),
]


def artefatos_do_documento():
    """Figuras que o documento efetivamente insere, lidas dos geradores."""
    figs = set()
    for g in ("gera_cap_v2.py", "gera_postextuais.py", "gera_pretextuais.py",
              "gera_introducao.py", "gera_metodologia.py"):
        f = GEN / g
        if f.exists():
            figs |= set(re.findall(r'"([a-z0-9_]+\.png)"',
                                   f.read_text(encoding="utf-8", errors="ignore")))
    return figs


def quem_produz(fig):
    """Script que grava a figura, incluindo os que gravam por f-string."""
    raiz = re.sub(r"_(ipca_geral|ipca_transportes|gasolina_refinaria|"
                  r"gasolina_consumidor|diesel|etanol)\.png$", "", fig)
    for f in sorted(COD.glob("*.py")):
        if f.name.startswith("auditoria"):
            continue
        s_ = f.read_text(encoding="utf-8", errors="ignore")
        if fig in s_ or (raiz != fig[:-4] and raiz + '_{' in s_) or \
           re.search(re.escape(raiz) + r'_?\{', s_):
            return f
    return None


def camada2():
    cab(2, "CADEIA DE AMOSTRA: figura do documento -> script -> planilha lida")
    print("   Esta e a camada que teria pego o bug de 05/09.")
    print("   Ela nao adivinha por nome de variavel: procura a extracao explicita")
    print("   de amostra cheia dentro do script que gera cada figura usada.\n")
    figs = artefatos_do_documento()
    print("   figuras inseridas no documento: %d\n" % len(figs))
    achou_alguma = False
    for fig in sorted(figs):
        src = quem_produz(fig)
        if src is None:
            aviso(2, "%s: nenhum script identificado como gerador" % fig)
            continue
        s_ = src.read_text(encoding="utf-8", errors="ignore")
        # olhar SO o corpo de figuras(), que e o caminho que produz a imagem;
        # o resto do arquivo pode legitimamente calcular amostra cheia para outros fins
        mf = re.search(chr(10) + r"def figuras\(\).*?(?=" + chr(10) + r"def |\Z)",
                       s_, re.S)
        alvo = mf.group(0) if mf else s_
        marcas = [rot for pad, rot in CHEIA if re.search(pad, alvo)]
        le = sorted(set(re.findall(r'["\']([a-z0-9_]+\.xlsx)["\']', s_)))
        linha = "   %-42s <- %-32s" % (fig, src.name)
        if marcas:
            achou_alguma = True
            print(linha + "  [AMOSTRA CHEIA]")
            crit(2, "%s vem de %s, que %s. Confirme se a figura usa a coluna de "
                    "regime ou a de amostra cheia." % (fig, src.name, "; ".join(marcas)))
            if le:
                print("               le: %s" % ", ".join(le)[:88])
        else:
            print(linha)
    if not achou_alguma:
        ok("nenhuma figura do documento vem de script que extraia amostra cheia")


# ══════════════════════════════════════════════════ 3 IDENTIDADES
def camada3():
    cab(3, "IDENTIDADES: as relacoes contabeis fecham?")
    # voltas por regime
    f = RES / "onda2_7b_voltas_por_regime.xlsx"
    if f.exists():
        d = pd.read_excel(f, sheet_name="decomposicao")
        e = (d.efeito_total - d.primeira_volta - d.segunda_volta).abs().max()
        (ok if e < 1e-8 else crit)("identidade total = 1a + 2a (por regime): desvio %.2e" % e) \
            if e < 1e-8 else crit(3, "identidade total = 1a + 2a NAO fecha: %.2e" % e)
        dif = d.dif_total.max() if "dif_total" in d else np.nan
        if dif == dif and dif > 0.05:
            crit(3, "total sob desenho comum difere da especificacao de referencia em %.4f" % dif)
        else:
            ok("total sob desenho comum bate com a referencia (desvio max %.4f)" % dif)
    else:
        crit(3, "ausente: %s" % f.name)
    # canal
    f = RES / "onda2_4_canal_combustiveis.xlsx"
    if f.exists():
        d = pd.read_excel(f, sheet_name="por_regime")
        if {"efeito_total", "efeito_via_combustivel", "efeito_nao_via_combustivel"} <= set(d.columns):
            e = (d.efeito_total - d.efeito_via_combustivel
                 - d.efeito_nao_via_combustivel).abs().max()
            (ok if e < 1e-6 else lambda m: crit(3, m))(
                "identidade total = canal + direto: desvio %.2e" % e)
        if "desvio_identidade" in d.columns:
            e2 = d.desvio_identidade.max()
            ok("FWL: canal = soma delta*b, desvio max %.2e" % e2) if e2 < 1e-6 \
                else crit(3, "FWL nao fecha: %.2e" % e2)


# ══════════════════════════════════════════════════ 4 CRUZADA
CRUZAR = [
    ("IPCA Geral, efeito total, regime posterior, h=0",
     [(RES / "onda3_3_regimes_data_escolhida.xlsx", None,
       lambda d: d[(d.variavel == "ipca_geral") & (d.h == 0)
                   & (d.data_corte.astype(str).str.startswith("2017-03"))].coef_pos.iloc[0]),
      (RES / "onda2_7b_voltas_por_regime.xlsx", "decomposicao",
       lambda d: d[(d.regime == "posterior") & (d.h == 0)
                   & (d.cenario == "central")].efeito_total.iloc[0])], 0.02),
    ("IPCA Geral, efeito total, regime posterior, h=3",
     [(RES / "onda3_3_regimes_data_escolhida.xlsx", None,
       lambda d: d[(d.variavel == "ipca_geral") & (d.h == 3)
                   & (d.data_corte.astype(str).str.startswith("2017-03"))].coef_pos.iloc[0]),
      (RES / "onda2_7b_voltas_por_regime.xlsx", "decomposicao",
       lambda d: d[(d.regime == "posterior") & (d.h == 3)
                   & (d.cenario == "central")].efeito_total.iloc[0])], 0.02),
]


def camada4():
    cab(4, "CRUZADA: a mesma grandeza tem o mesmo valor em toda parte?")
    for rot, fontes, tol in CRUZAR:
        vs = []
        for f, ab, fn in fontes:
            if not f.exists():
                crit(4, "%s: falta %s" % (rot, f.name)); vs = None; break
            d = pd.read_excel(f, sheet_name=ab) if ab else pd.read_excel(f)
            try:
                vs.append((f.name, float(fn(d))))
            except Exception as e:
                crit(4, "%s: erro lendo %s (%s)" % (rot, f.name, e)); vs = None; break
        if not vs:
            continue
        span = max(v for _, v in vs) - min(v for _, v in vs)
        txt = " vs ".join("%.3f (%s)" % (v, n.split("_")[0]) for n, v in vs)
        if span > tol:
            crit(4, "%s: %s  -> divergencia %.3f > %.3f" % (rot, txt, span, tol))
        else:
            ok("%s: %s  (divergencia %.3f)" % (rot, txt, span))


# ══════════════════════════════════════════════════ 5 CADEIA
def camada5():
    cab(5, "CADEIA: dado -> script -> resultado -> figura -> documento")
    def mt(p): return os.path.getmtime(p) if Path(p).exists() else None
    dado = mt(ART / "base/base_corrigida.parquet")
    espec = mt(COD / "lp_base_nova.py")
    doc = mt(DOCX)
    base = max(x for x in (dado, espec) if x)
    print("   dado    %s" % time.strftime("%d/%m %H:%M", time.localtime(dado)))
    print("   espec   %s" % time.strftime("%d/%m %H:%M", time.localtime(espec)))
    print("   docx    %s" % time.strftime("%d/%m %H:%M", time.localtime(doc)))
    # Deliberadamente congelados: rodam no corte antigo e sem dummies de pandemia,
    # de proposito, porque sao comparacao com o estado anterior. O proprio
    # lp_base_nova.py adverte "Nao tirar numero de la".
    CONGELADOS = {"aux_comparacao_juros_ols.xlsx", "aux_comparacao_juros_iv.xlsx",
                  "aux_comparacao_juros_regimes.xlsx", "aux_comparacao_juros_resumo.xlsx"}
    usadas = artefatos_do_documento()
    velhos, ignorados, congelados = [], 0, 0
    for f in relevantes(RES, "*.xlsx") + list(FIG.glob("*.png")):
        m = mt(f)
        if not m or m >= base:
            continue
        if f.name in CONGELADOS:
            congelados += 1
            continue
        if f.suffix == ".png" and f.name not in usadas:
            ignorados += 1                       # figura antiga que o TCC nao usa
            continue
        velhos.append((f.name, m))
    if ignorados:
        print("   (%d figura(s) antiga(s) ignorada(s) por nao entrarem no documento)"
              % ignorados)
    if congelados:
        print("   (%d arquivo(s) de comparacao deliberadamente congelado(s), na lista branca)"
              % congelados)
    if velhos:
        for n, m in sorted(velhos, key=lambda x: x[1]):
            crit(5, "%s ENTRA NO DOCUMENTO e e anterior ao dado/especificacao (%s)"
                    % (n, time.strftime("%d/%m %H:%M", time.localtime(m))))
    else:
        ok("nenhum artefato usado no documento e anterior ao dado ou a especificacao")
    if doc < max([mt(f) for f in FIG.glob("*.png")] + [base]):
        crit(5, "o .docx e anterior a alguma figura: reconstrua o documento")
    else:
        ok("o documento e o ultimo elo da cadeia")


# ══════════════════════════════════════════════════ 6 REMISSOES
def camada6(paras, cels):
    cab(6, "REMISSOES E LISTAS")
    txt = " ".join(paras) + " " + " ".join(" ".join(r) for r in cels)
    i = next((k for k, p in enumerate(paras)
              if p.strip().upper().startswith("1 INTRODU")), 0)
    pre, corpo = " ".join(paras[:i]), " ".join(paras[i:])

    def rot(t, k):
        return {m.group(1) for m in re.finditer(k + r" ((?:\d+|[A-Z]\.\d+)) –", t)}
    for k in ("Figura", "Tabela"):
        L, C = rot(pre, k), rot(corpo, k)
        if L == C:
            ok("lista de %ss bate com o conteudo (%d itens)" % (k.lower(), len(L)))
        else:
            crit(6, "%ss so na lista: %s | so no texto: %s"
                    % (k, sorted(L - C) or "-", sorted(C - L) or "-"))
    secs = {m.group(1) for m in re.finditer(r"^(\d(?:\.\d+)*)\s+\S", "\n".join(paras), re.M)}
    for m in re.finditer(r"seç(?:ão|oes)\s+(\d(?:\.\d+)*)", txt):
        if m.group(1) not in secs:
            crit(6, "remissao a secao %s, que nao existe" % m.group(1))
    aps = {m.group(1) for m in re.finditer(r"APÊNDICE ([A-Z])", txt.upper())}
    for m in re.finditer(r"Apêndice ([A-Z])\b", txt):
        if m.group(1) not in aps:
            crit(6, "remissao ao Apendice %s, que nao existe" % m.group(1))
    ok("remissoes a secao e apendice verificadas")
    if any(chr(8212) in p for p in paras):
        crit(6, "ha travessao no texto")
    else:
        ok("nenhum travessao")


# ══════════════════════════════════════════════════
def main():
    so = sys.argv[1] if len(sys.argv) > 1 else None
    print("AUDITORIA 360  |  %s" % time.strftime("%d/%m/%Y %H:%M"))
    print("documento: %s" % DOCX.name)
    try:
        _, paras, cels = ler_docx()
    except Exception as e:
        print("\nNAO FOI POSSIVEL ABRIR O DOCX (%s)." % type(e).__name__)
        print("Se estiver aberto no Word, feche e rode de novo.")
        raise SystemExit(2)
    for n, fn in [("1", lambda: camada1(paras, cels)), ("2", camada2),
                  ("3", camada3), ("4", camada4), ("5", camada5),
                  ("6", lambda: camada6(paras, cels))]:
        if so in (None, n):
            fn()
    print("\n" + "=" * 78)
    print("RESUMO: %d achado(s) critico(s), %d aviso(s)" % (len(CRIT), len(AVISO)))
    for c, m in CRIT:
        print("   CRITICO camada %s: %s" % (c, m))
    print("=" * 78)
    raise SystemExit(1 if CRIT else 0)


if __name__ == "__main__":
    main()
