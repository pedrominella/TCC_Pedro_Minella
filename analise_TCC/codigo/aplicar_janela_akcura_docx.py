# -*- coding: utf-8 -*-
"""
Troca a janela da comparacao internacional (secao 4.8, Apendice I, Figura 13) de
ago/2020-jun/2022 para abr/2020-jun/2022, o segundo subperiodo de Akcura (2025),
que vai do vale do Brent em abril de 2020 ao pico de junho de 2022 (08/10/2026).
Numeros de fig_repasse_internacional.carrega(), conferidos com assert antes de gravar.
Backup feito antes: historico/3_versoes_e_backups/backups_do_TCC/_backup_TCC_pre_janela_akcura.docx
"""
import io, sys
from pathlib import Path
import docx
import pandas as pd
from PIL import Image

AQUI = Path(__file__).resolve().parent
sys.path.insert(0, str(AQUI))
import fig_repasse_internacional as FR

RAIZ = AQUI.parents[1]
ARQ = RAIZ / "TCC_PEDRO_MINELLA.docx"
FIG = RAIZ / "analise_TCC" / "figuras" / "repasse_internacional.png"


def v3(x):
    return ("%.3f" % x).replace(".", ",")


# ---------------------------------------------------------------- numeros
assert FR.A == pd.Period("2020-04", "M") and FR.B == pd.Period("2022-06", "M")
t = FR.carrega()
c = pd.ExcelFile(FR.BASE / "Global_Fuel_Subsidies_&_Price_Control_Measures_Database.xlsx")
fundo = {}
for ano in ["2021", "2022"]:
    for _, r in c.parse(ano, header=None, skiprows=2).iterrows():
        fundo.setdefault(str(r[1]).strip(), []).append(r[9])  # coluna "Fuel Price Stabilization Fund or Mechanism"
t["fundo"] = t.cod.map(lambda k: "sim" if any(str(v) == "1" for v in fundo.get(k, [])) else "não")

br = float(t.loc[t.cod == "BRA", "pt"].iloc[0])
ordem = t.sort_values("pt", ascending=False).reset_index(drop=True)
pos = int(ordem.index[ordem.cod == "BRA"][0]) + 1
pct = 100 * (1 - (pos - 1) / (len(t) - 1))
grupos = {"desregulado": t[t.regime == "desregulado"], "regulado": t[t.regime == "regulado"], "todas": t}
med = {k: s.pt.median() for k, s in grupos.items()}
lac = t[t.cod.isin(FR.LAC)].sort_values("pt", ascending=False).reset_index(drop=True)
pos_lac = int(lac.index[lac.cod == "BRA"][0]) + 1

assert len(t) == 65 and len(grupos["desregulado"]) == 16 and len(grupos["regulado"]) == 49
assert v3(br) == "0,868" and pos == 16 and round(pct) == 77
assert v3(med["desregulado"]) == "0,798" and v3(med["regulado"]) == "0,288"
assert round(med["desregulado"] / med["regulado"], 1) == 2.8
assert len(lac) == 16 and pos_lac == 6
assert list(lac.cod[:3]) == ["PAN", "BLZ", "PER"] and list(lac.regime[:3]) == ["regulado", "regulado", "desregulado"]
assert lac.loc[2, "fundo"] == "sim" and lac.loc[lac.cod == "BRA", "fundo"].iloc[0] == "não"
q3_desreg = grupos["desregulado"].pt.quantile(0.75)
assert med["desregulado"] < br < q3_desreg

# ---------------------------------------------------------------- documento
d = docx.Document(ARQ)
P = d.paragraphs


def trocar_tudo(p, novo):
    """Texto inteiro no primeiro run (mantem a formatacao dele), demais runs esvaziados."""
    p.runs[0].text = novo
    for r in p.runs[1:]:
        r.text = ""


def trocar_trecho(p, velho, novo):
    for r in p.runs:
        if velho in r.text:
            r.text = r.text.replace(velho, novo)
            return
    raise ValueError(f"trecho nao encontrado em um unico run: {velho!r}")


def conferir(i, inicio):
    assert P[i].text.startswith(inicio), (i, P[i].text[:80])


for i in (124, 561):
    conferir(i, "Tabela I.1 – ")
    trocar_trecho(P[i], "agosto de 2020", "abril de 2020")

conferir(415, "O contraste documentado neste capítulo")
trocar_trecho(P[415], "A janela é a alta de agosto de 2020 a junho de 2022.",
              "A janela é a alta de abril de 2020 a junho de 2022, o segundo subperíodo daquele trabalho, "
              "que vai do menor preço do Brent em abril de 2020 ao maior em junho de 2022.")

conferir(416, "O resultado consta da Figura 13")
trocar_tudo(P[416], (
    "O resultado consta da Figura 13, no Apêndice N. Entre os 65 países com série utilizável, aqueles que a "
    "base do Banco Mundial (WORLD BANK, 2025) classifica como de preço desregulado apresentam repasse mediano "
    f"de {v3(med['desregulado'])}, contra {v3(med['regulado'])} nos de preço regulado, razão de cerca de 2,8 "
    "vezes que reproduz, em corte transversal e com dados independentes, o mecanismo que este trabalho estima "
    f"para o Brasil ao longo do tempo. O Brasil aparece com {v3(br)}, no percentil {pct:.0f} da distribuição e "
    "em sexto lugar entre os 16 países da América Latina e Caribe. A associação não é mecânica: os dois maiores "
    "repasses da região são os do Panamá e de Belize, classificados como de preço regulado, e o Peru, que mantém "
    "fundo de estabilização, vem em terceiro, o que sugere que importa menos a existência do controle ou do "
    "fundo do que a regra pela qual eles são acionados."))

conferir(427, "A literatura de corte transversal")
trocar_trecho(P[427], "que os países de preço desregulado repassam duas vezes e meia mais que os de preço regulado",
              "que o repasse mediano dos países de preço desregulado é cerca de 2,8 vezes o dos países de preço regulado")

conferir(560, "Este apêndice reporta os números")
trocar_trecho(P[560], "entre agosto de 2020 e junho de 2022",
              "entre abril de 2020 e junho de 2022, o segundo subperíodo daquele trabalho")

conferir(564, "A ordenação regional não se explica")
trocar_tudo(P[564], (
    "A ordenação regional não se explica apenas pela classificação binária da regra de preço. Panamá e Belize, "
    "classificados como de preço regulado, apresentam os dois maiores repasses da região, e o Peru, classificado "
    "como desregulado, vem em terceiro ainda mantendo fundo de estabilização. O que o quadro sustenta é a "
    "associação central, e não uma correspondência termo a termo: a Tabela I.2 mostra que a mediana dos países "
    "de preço desregulado é cerca de 2,8 vezes a dos países de preço regulado, e o Brasil situa-se acima da "
    "mediana do próprio grupo a que pertence, entre ela e o terceiro quartil."))

conferir(566, "Nota: o Brasil ocupa a oitava posição")
trocar_tudo(P[566], (
    f"Nota: o Brasil ocupa a décima sexta posição entre as 65 economias, o que corresponde ao percentil {pct:.0f} "
    "da distribuição. A amostra restringe-se às economias com observação de preço de varejo em ambos os extremos "
    "da janela."))

conferir(659, "Nota: repasse calculado pela medida")
trocar_trecho(P[659], "na alta de agosto de 2020 a junho de 2022", "na alta de abril de 2020 a junho de 2022")

# Tabela I.1: America Latina e Caribe, ordenada
ti = d.tables[7]
assert ti.rows[0].cells[1].text == "País" and len(ti.rows) == 17
for k, r in lac.iterrows():
    vals = [str(k + 1), FR._PAIS_PT.get(r.cod, r.pais), v3(r.pt), r.regime, r.fundo]
    for cel, v in zip(ti.rows[k + 1].cells, vals):
        trocar_tudo(cel.paragraphs[0], v)

# Tabela I.2: distribuicao por regra de preco
tq = d.tables[8]
assert tq.rows[0].cells[2].text == "1º quartil" and tq.rows[4].cells[0].text == "Brasil"
for linha, chave in ((1, "desregulado"), (2, "regulado"), (3, "todas")):
    s = grupos[chave].pt
    vals = [str(len(s)), v3(s.quantile(0.25)), v3(s.median()), v3(s.quantile(0.75)), v3(s.mean())]
    for cel, v in zip(tq.rows[linha].cells[1:], vals):
        trocar_tudo(cel.paragraphs[0], v)
trocar_tudo(tq.rows[4].cells[3].paragraphs[0], v3(br))

# Figura 13: troca a imagem e ajusta a altura mantendo a largura
blip = P[658]._p.xpath(".//a:blip/@r:embed")[0]
parte = d.part.related_parts[blip]
novo = FIG.read_bytes()
parte._blob = novo
w, h = Image.open(io.BytesIO(novo)).size
for ext in P[658]._p.xpath(".//wp:extent") + P[658]._p.xpath(".//a:ext"):
    cx = int(ext.get("cx"))
    ext.set("cy", str(int(cx * h / w)))

d.save(ARQ)
print("salvo", ARQ)
print("Brasil %s | %d de 65 | percentil %.0f | %dº na AL | medianas %s e %s"
      % (v3(br), pos, pct, pos_lac, v3(med["desregulado"]), v3(med["regulado"])))
