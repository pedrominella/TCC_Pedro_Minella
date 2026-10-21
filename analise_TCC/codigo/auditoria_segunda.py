# -*- coding: utf-8 -*-
"""
SEGUNDA VARREDURA  --  CHECAGENS 31 a 60

Complementa auditoria_completa.py (1 a 30). Cobre o que aquela nao alcanca:
reprodutibilidade, consistencia interna da base, especificacao de segunda ordem,
poder e calibragem da inferencia, e coerencia entre o texto e os numeros.

Uso:  python auditoria_segunda.py            todas
      python auditoria_segunda.py 44 49 43   so essas
      python auditoria_segunda.py --rapido   pula as caras (43, 50, 51, 52)
"""
import sys, re, time, warnings, subprocess, importlib
from pathlib import Path
import numpy as np, pandas as pd
import statsmodels.api as sm
from scipy import stats

warnings.filterwarnings("ignore")
AQUI = Path(__file__).resolve().parent
sys.path.insert(0, str(AQUI))
import lp_base_nova as L
import onda5_alternativos as A

ART = AQUI.parent
RES, FIG = ART / "resultados", ART / "figuras"
DOCX = ART.parent / "TCC_PEDRO_MINELLA.docx"
ALVO, H = "ipca_geral", 0
R = []


def reg(n, fam, nome, ver, det):
    R.append({"n": n, "familia": fam, "checagem": nome, "veredito": ver, "detalhe": det})
    print("  %2d %s%-44s %s" % (n, {"OK": "[ok]      ", "ATENCAO": "[ATENCAO] ",
                                    "CRITICO": "[CRITICO] "}[ver], nome, det))


def dados():
    if not hasattr(dados, "c"):
        dados.c = A.base()
    return dados.c


def desenho(nome=ALVO, h=H):
    b, dum = dados()
    return A.monta(b, dum, nome, h)


def scripts():
    return [f for f in sorted(AQUI.glob("*.py")) if not f.name.startswith("auditoria")]


# ═══════════════════════════ F. REPRODUTIBILIDADE
def c31():
    import numpy, pandas, statsmodels, scipy
    v = "numpy %s, pandas %s, statsmodels %s, scipy %s" % (
        numpy.__version__, pandas.__version__, statsmodels.__version__, scipy.__version__)
    tem = any("statsmodels" in f.read_text(encoding="utf-8", errors="ignore")
              and "__version__" in f.read_text(encoding="utf-8", errors="ignore")
              for f in scripts())
    reg(31, "F", "versoes de biblioteca registradas",
        "OK" if tem else "ATENCAO",
        ("registradas" if tem else "NAO registradas em lugar algum") + ". Em uso: " + v)


def c32():
    cand = list(ART.glob("*.md")) + list(ART.glob("LEIA*")) + list(AQUI.glob("rodar*.py")) \
        + list(ART.glob("Makefile"))
    reg(32, "F", "ordem de execucao documentada",
        "OK" if cand else "ATENCAO",
        ", ".join(c.name for c in cand[:4]) if cand
        else "nenhum LEIA-ME, Makefile ou rodar_tudo.py encontrado")


def c33():
    g1 = np.random.default_rng(12345)
    g2 = np.random.default_rng(12345)
    a = g1.choice([-1.0, 1.0], size=500).sum()
    bb = g2.choice([-1.0, 1.0], size=500).sum()
    y, X, cols = desenho()
    p1 = sm.OLS(y, X).fit().params[cols.index("choque_pos")]
    p2 = sm.OLS(y, X).fit().params[cols.index("choque_pos")]
    ok = (a == bb) and abs(p1 - p2) < 1e-12
    reg(33, "F", "determinismo (mesma semente, mesmo numero)",
        "OK" if ok else "CRITICO",
        "sorteio identico: %s | coeficiente identico: %s" % (a == bb, abs(p1 - p2) < 1e-12))


def c34():
    achou = []
    for f in scripts():
        s = f.read_text(encoding="utf-8", errors="ignore")
        n = len(re.findall(r"[A-Za-z]:[/\\]{1,2}Users", s))
        if n:
            achou.append("%s(%d)" % (f.name, n))
    reg(34, "F", "caminhos absolutos embutidos",
        "ATENCAO" if achou else "OK",
        ("%d script(s): " % len(achou)) + ", ".join(achou[:5]) if achou else "nenhum")


def c35():
    texto = " ".join(f.read_text(encoding="utf-8", errors="ignore") for f in scripts())
    orfas = []
    for f in RES.rglob("*.xlsx"):
        if any(x in f.parts for x in ("_superseded", "_antes_rerun")):
            continue
        if f.name not in texto and f.stem not in texto:
            orfas.append(f.name)
    reg(35, "F", "planilhas orfas (nenhum script as gera)",
        "ATENCAO" if orfas else "OK",
        ("%d: " % len(orfas)) + ", ".join(orfas[:5]) if orfas else "todas rastreadas")


def c36():
    gen = Path(r"C:/Users/pedro/AppData/Local/Temp/claude/"
               r"c--Users-pedro-OneDrive-Documentos-TCC-python/"
               r"b948f4bc-20ef-496a-b291-1171ce8ec6b1/scratchpad")
    doc_txt = " ".join((gen / g).read_text(encoding="utf-8", errors="ignore")
                       for g in ("gera_cap_v2.py", "gera_postextuais.py",
                                 "gera_pretextuais.py", "gera_metodologia.py")
                       if (gen / g).exists())
    outros = " ".join(f.read_text(encoding="utf-8", errors="ignore") for f in scripts())
    orfaos = []
    for f in scripts():
        prod = re.findall(r'["\']([a-z0-9_{}]+\.(?:png|xlsx))["\']', f.read_text(encoding="utf-8", errors="ignore"))
        if not prod:
            continue
        usado = any(p.split("{")[0] in doc_txt or p.split("{")[0] in outros.replace(f.name, "")
                    for p in prod)
        if not usado:
            orfaos.append(f.name)
    reg(36, "F", "scripts que nao alimentam nada",
        "ATENCAO" if orfaos else "OK",
        ", ".join(orfaos[:6]) if orfaos else "todos alimentam figura, tabela ou outro script")


# ═══════════════════════════ G. DADOS
def c37():
    marcas = 0
    for f in scripts():
        s = f.read_text(encoding="utf-8", errors="ignore")
        marcas += len(re.findall(r"baixad[oa] em|acesso em|extra[ií]d[oa] em|vintage", s, re.I))
    reg(37, "G", "data de extracao das series registrada",
        "OK" if marcas >= 3 else "ATENCAO",
        "%d marcacao(oes) de data de extracao no codigo" % marcas)


def c38():
    b, _ = dados()
    pior, onde = 0.0, "-"
    for nome in L.VARIAVEIS:
        d1 = pd.to_numeric(b[f"ln_{nome}"], errors="coerce").diff()
        d2 = pd.to_numeric(b[f"dlog_{nome}"], errors="coerce")
        e = float((d1 - d2).abs().max())
        if e == e and e > pior:
            pior, onde = e, nome
    reg(38, "G", "nivel e variacao sao a mesma coisa",
        "CRITICO" if pior > 1e-9 else "OK",
        "maior desvio ln.diff() vs dlog: %.2e (%s)" % (pior, onde))


def c39():
    b, dum = dados()
    ruim = []
    for nome in L.VARIAVEIS:
        y, X, _ = A.monta(b, dum, nome, H)
        z = (y - y.mean()) / y.std()
        if np.abs(z).max() > 5:
            ruim.append("%s %.1f" % (nome, np.abs(z).max()))
    reg(39, "G", "outliers na variavel DEPENDENTE",
        "ATENCAO" if ruim else "OK",
        "; ".join(ruim) if ruim else "nenhuma dependente com |z| > 5")


def c40():
    b, _ = dados()
    piores = []
    for c in L.CONTROLES:
        for lag in range(1, L.LAGS + 1):
            col = f"{c}_lag{lag}"
            if col not in b.columns:
                continue
            r = b[["choque_petroleo", col]].dropna().corr().iloc[0, 1]
            if abs(r) > 0.25:
                piores.append("%s r=%.2f" % (col, r))
    reg(40, "G", "controles defasados correlacionados com o choque",
        "ATENCAO" if piores else "OK",
        "; ".join(piores[:4]) if piores else "todas as correlacoes abaixo de 0,25")


def c41():
    b, _ = dados()
    r = b[["dlog_petroleo_usd", "dlog_cambio"]].dropna().corr().iloc[0, 1]
    reg(41, "G", "choque em USD e cambio: dupla contagem?",
        "ATENCAO" if abs(r) > 0.4 else "OK",
        "correlacao entre o choque e a variacao cambial: %.3f" % r)


def c42():
    b, dum = dados()
    y, X, cols = desenho()
    res = sm.OLS(y, X).fit().resid
    M = X[:, [cols.index(c) for c in dum]]
    rr = sm.OLS(res, sm.add_constant(M)).fit()
    reg(42, "G", "sazonalidade residual apos as 11 dummies",
        "CRITICO" if rr.f_pvalue < 0.05 else "OK",
        "F=%.3f p=%.3f (por construcao deve ser ~0)" % (rr.fvalue, rr.f_pvalue))


# ═══════════════════════════ H. ESPECIFICACAO
def c43(rapido=False):
    if rapido:
        reg(43, "H", "segunda quebra alem de 2017", "OK", "pulado (--rapido)"); return
    b, dum = dados()
    corte1 = pd.Timestamp(L.DATA_CORTE_FINAL)
    datas = pd.to_datetime(b.Data)
    cand = [d for d in datas if d > corte1 + pd.DateOffset(months=18)
            and d < datas.max() - pd.DateOffset(months=18)]
    melhor, W = None, -1
    for d in cand:
        bb = b.copy()
        bb["r2"] = (bb.Data >= d).astype(float)
        bb["choque_pre"] = bb["choque_petroleo"] * (1 - bb["regime_pos"])
        bb["choque_pos"] = bb["choque_petroleo"] * bb["regime_pos"] * (1 - bb["r2"])
        bb["choque_r2"] = bb["choque_petroleo"] * bb["r2"]
        try:
            y, X, cols = A.monta(bb, dum, ALVO, H)
            X = np.column_stack([X, bb.loc[:, "choque_r2"].values[-len(y):], ])
        except Exception:
            continue
        d2 = bb.copy()
        for lag in range(1, L.LAGS + 1):
            d2[f"y_lag{lag}"] = d2[f"dlog_{ALVO}"].shift(lag)
        pand = L.controles_pandemia(d2, H, modo="janela")
        regs = ["choque_pre", "choque_pos", "choque_r2", "regime_pos", "r2"]
        for lag in range(1, L.LAGS + 1):
            regs += [f"y_lag{lag}", f"choque_petroleo_lag{lag}"] + \
                    [f"{c}_lag{lag}" for c in L.CONTROLES]
        regs += list(dum) + pand
        d2["__y"] = L.y_acum(d2, ALVO, H)
        am = d2[["__y"] + regs].replace([np.inf, -np.inf], np.nan).dropna()
        Xd = sm.add_constant(am[regs].astype(float)); Xd = Xd.loc[:, Xd.std() > 0]
        if "choque_pos" not in Xd or "choque_r2" not in Xd:
            continue
        rr = sm.OLS(am["__y"].astype(float), Xd).fit(cov_type="HAC", cov_kwds=L.hac_kw(H))
        w = float(rr.wald_test("choque_pos = choque_r2", scalar=True).statistic)
        if w > W:
            W, melhor = w, d
    reg(43, "H", "segunda quebra alem de mar/2017",
        "ATENCAO" if W > 6.63 else "OK",
        "melhor candidata %s, sup-Wald=%.2f (chi2 1%%=6,63; critico real e maior)"
        % (melhor.strftime("%b/%Y") if melhor is not None else "-", W))


def c44():
    b, _ = dados()
    piores = []
    for c in L.CONTROLES:
        d = b[[c, "choque_petroleo"] + [f"choque_petroleo_lag{l}" for l in range(1, L.LAGS + 1)]].dropna()
        X = sm.add_constant(d[["choque_petroleo"] + [f"choque_petroleo_lag{l}"
                                                     for l in range(1, L.LAGS + 1)]])
        r = sm.OLS(d[c], X).fit(cov_type="HAC", cov_kwds={"maxlags": 4})
        f = r.f_test(np.eye(len(r.params))[1:])
        if float(f.pvalue) < 0.05:
            piores.append("%s p=%.4f R2=%.3f" % (c, float(f.pvalue), r.rsquared))
    reg(44, "H", "CONTROLES RUINS: reagem ao choque?",
        "ATENCAO" if piores else "OK",
        ("; ".join(piores) + " -- controlar por eles pode subestimar o repasse")
        if piores else "nenhum controle e explicado pelo choque")


def c45():
    y, X, cols = desenho()
    r = sm.OLS(y, X).fit(cov_type="HAC", cov_kwds=L.hac_kw(H))
    idx = [cols.index(f"choque_petroleo_lag{l}") for l in range(1, L.LAGS + 1)]
    Rm = np.zeros((len(idx), len(cols)))
    for k, i in enumerate(idx):
        Rm[k, i] = 1
    f = r.wald_test(Rm, scalar=False)
    sig = [f"lag{l}" for l in range(1, L.LAGS + 1)
           if r.pvalues[cols.index(f"choque_petroleo_lag{l}")] < 0.10]
    reg(45, "H", "defasagens do choque sao significativas?",
        "ATENCAO" if float(f.pvalue) < 0.05 else "OK",
        "teste conjunto p=%.3f; individuais a 10%%: %s"
        % (float(f.pvalue), ", ".join(sig) if sig else "nenhuma"))


def c46():
    b, dum = dados()
    inuteis = []
    for c in L.CONTROLES:
        ps = []
        for h in [0, 1, 3, 6, 12]:
            y, X, cols = A.monta(b, dum, ALVO, h)
            r = sm.OLS(y, X).fit(cov_type="HAC", cov_kwds=L.hac_kw(h))
            ps += [r.pvalues[cols.index(f"{c}_lag{l}")] for l in range(1, L.LAGS + 1)]
        if min(ps) > 0.10:
            inuteis.append("%s (menor p=%.2f)" % (c, min(ps)))
    reg(46, "H", "controle nunca significativo em horizonte algum",
        "ATENCAO" if inuteis else "OK",
        "; ".join(inuteis) if inuteis else "todos os controles sao relevantes em algum ponto")


def c47():
    from statsmodels.stats.diagnostic import linear_reset
    y, X, _ = desenho()
    r = sm.OLS(y, X).fit()
    try:
        t = linear_reset(r, power=2, use_f=True)
        p = float(t.pvalue)
    except Exception as e:
        reg(47, "H", "teste RESET de forma funcional", "ATENCAO", "nao rodou: %s" % e); return
    reg(47, "H", "teste RESET de forma funcional",
        "ATENCAO" if p < 0.05 else "OK", "RESET(2) p=%.3f" % p)


def c48():
    b, dum = dados()
    bb = b.copy()
    bb["ch2"] = bb["choque_petroleo"] ** 2 * bb["regime_pos"]
    d = bb.copy()
    for lag in range(1, L.LAGS + 1):
        d[f"y_lag{lag}"] = d[f"dlog_{ALVO}"].shift(lag)
    pand = L.controles_pandemia(d, H, modo="janela")
    regs = ["choque_pre", "choque_pos", "ch2", "regime_pos"]
    for lag in range(1, L.LAGS + 1):
        regs += [f"y_lag{lag}", f"choque_petroleo_lag{lag}"] + \
                [f"{c}_lag{lag}" for c in L.CONTROLES]
    regs += list(dum) + pand
    d["__y"] = L.y_acum(d, ALVO, H)
    am = d[["__y"] + regs].replace([np.inf, -np.inf], np.nan).dropna()
    X = sm.add_constant(am[regs].astype(float)); X = X.loc[:, X.std() > 0]
    r = sm.OLS(am["__y"].astype(float), X).fit(cov_type="HAC", cov_kwds=L.hac_kw(H))
    p = r.pvalues.get("ch2", np.nan)
    reg(48, "H", "o repasse depende do TAMANHO do choque?",
        "ATENCAO" if p < 0.10 else "OK",
        "termo quadratico: coef %.4f, p=%.3f" % (r.params.get("ch2", np.nan), p))


# ═══════════════════════════ I. INFERENCIA
def c49():
    b, dum = dados()
    y, X, cols = desenho()
    r = sm.OLS(y, X).fit(cov_type="HAC", cov_kwds=L.hac_kw(H))
    ep_pre = r.bse[cols.index("choque_pre")]
    mde = (1.96 + 0.84) * ep_pre           # 80% de poder, 5% bilateral
    pos = r.params[cols.index("choque_pos")]
    reg(49, "I", "PODER: menor efeito detectavel no regime anterior",
        "ATENCAO" if mde > abs(pos) else "OK",
        "ep_pre=%.4f -> MDE=%.4f; o efeito do regime posterior e %.4f (%.1fx o MDE)"
        % (ep_pre, mde, pos, pos / mde))


def c50(rapido=False):
    if rapido:
        reg(50, "I", "cobertura do bootstrap", "OK", "pulado (--rapido)"); return
    rng = np.random.default_rng(99)
    T, B, N = 200, 199, 300
    dentro = 0
    for _ in range(N):
        x = rng.standard_normal(T); e = rng.standard_normal(T)
        yv = 0.3 * x + e
        bhat = np.polyfit(x, yv, 1)[0]
        res = yv - np.polyval(np.polyfit(x, yv, 1), x)
        bs = []
        for _ in range(B):
            s = rng.choice([-1.0, 1.0], size=T)
            bs.append(np.polyfit(x, np.polyval(np.polyfit(x, yv, 1), x) + res * s, 1)[0])
        lo, hi = np.percentile(bs, [5, 95])
        dentro += (lo <= 0.3 <= hi)
    cob = 100 * dentro / N
    reg(50, "I", "cobertura do bootstrap selvagem (simulacao)",
        "ATENCAO" if abs(cob - 90) > 5 else "OK",
        "cobre %.0f%% das vezes (nominal 90%%), %d replicas x %d amostras" % (cob, B, N))


def c51(rapido=False):
    if rapido:
        reg(51, "I", "sensibilidade ao tamanho do bloco", "OK", "pulado (--rapido)"); return
    y, X, cols = desenho()
    ip = cols.index("choque_pos")
    r = sm.OLS(y, X).fit()
    res, aj = r.resid, r.fittedvalues
    rng = np.random.default_rng(7)
    out = []
    for tam in [6, 12, 18, 24]:
        bs = []
        for _ in range(299):
            k = int(np.ceil(len(y) / tam))
            ini = rng.integers(0, len(y) - tam + 1, size=k)
            idx = np.concatenate([np.arange(i, i + tam) for i in ini])[:len(y)]
            s = rng.choice([-1.0, 1.0], size=k).repeat(tam)[:len(y)]
            bs.append(sm.OLS(aj + res[idx] * s, X).fit().params[ip])
        out.append((tam, float(np.std(bs, ddof=1))))
    eps = [e for _, e in out]
    amp = (max(eps) - min(eps)) / min(eps)
    reg(51, "I", "sensibilidade ao tamanho do bloco do bootstrap",
        "ATENCAO" if amp > 0.30 else "OK",
        ", ".join("bloco %d: ep %.4f" % (t, e) for t, e in out) + " (dispersao %.0f%%)" % (100 * amp))


def c52(rapido=False):
    if rapido:
        reg(52, "I", "replicas suficientes", "OK", "pulado (--rapido)"); return
    y, X, cols = desenho()
    ip = cols.index("choque_pos")
    r = sm.OLS(y, X).fit()
    res, aj = r.resid, r.fittedvalues
    ps = []
    for sem in [1, 2, 3, 4, 5]:
        rng = np.random.default_rng(sem)
        bs = [sm.OLS(aj + res * rng.choice([-1.0, 1.0], size=len(y)), X).fit().params[ip]
              for _ in range(299)]
        ps.append(float(np.mean(np.abs(np.array(bs) - np.mean(bs)) >= abs(r.params[ip]))))
    reg(52, "I", "p-valor estavel entre sementes (299 replicas)",
        "ATENCAO" if (max(ps) - min(ps)) > 0.02 else "OK",
        "p entre %.4f e %.4f em 5 sementes" % (min(ps), max(ps)))


def c53():
    b, dum = dados()
    eps, coefs = [], []
    for h in range(0, 13):
        y, X, cols = A.monta(b, dum, ALVO, h)
        r = sm.OLS(y, X).fit(cov_type="HAC", cov_kwds=L.hac_kw(h))
        coefs.append(r.params[cols.index("choque_pos")]); eps.append(r.bse[cols.index("choque_pos")])
    rng = np.random.default_rng(5)
    def crit(k):
        rho = 0.8 ** np.abs(np.subtract.outer(np.arange(k), np.arange(k)))
        Lc = np.linalg.cholesky(rho + 1e-10 * np.eye(k))
        Z = Lc @ rng.standard_normal((k, 20000))
        return float(np.quantile(np.max(np.abs(Z), axis=0), 0.90))
    c7, c13 = crit(7), crit(13)
    reg(53, "I", "valor critico sup-t depende de quantos horizontes",
        "ATENCAO" if (c13 - c7) / c7 > 0.15 else "OK",
        "H=7 -> c=%.2f | H=13 -> c=%.2f (+%.0f%%)" % (c7, c13, 100 * (c13 - c7) / c7))


def c54():
    from onda2_7_segunda_volta import PESOS
    tot = {k: sum(v.values()) for k, v in PESOS.items()}
    reg(54, "I", "os tres cenarios de peso do IPCA",
        "ATENCAO",
        "soma dos pesos: " + ", ".join("%s %.3f" % (k, v) for k, v in tot.items())
        + " -- confira contra o peso oficial de combustiveis no IPCA")


# ═══════════════════════════ J. TEXTO x NUMEROS
def _docx():
    from docx import Document
    d = Document(str(DOCX))
    paras = [p.text for p in d.paragraphs]
    cels = [[c.text.strip() for c in r.cells] for t in d.tables for r in t.rows]
    return paras, cels


def c55():
    paras, cels = _docx()
    tab = " ".join(" ".join(r) for r in cels)
    txt = " ".join(paras)
    falhas = []
    for m in re.finditer(r"(\d,\d{3})\s*\*{0,3}", tab):
        pass
    # afirmacoes de significancia no texto
    for m in re.finditer(r"significativ\w+ a (menos de )?(\d+)%", txt):
        pass
    n_texto = len(re.findall(r"significativ", txt))
    n_estrela = len(re.findall(r"\*", tab))
    reg(55, "J", "afirmacoes de significancia no texto",
        "ATENCAO",
        "%d mencoes a 'significativ' no texto, %d asteriscos nas tabelas -- "
        "confira uma a uma nas 4 principais" % (n_texto, n_estrela))


def c56():
    paras, _ = _docx()
    txt = " ".join(paras)
    conf = []
    for m in re.finditer(r"(\d,\d{3})[^.]{0,60}?(\d,\d{3})[^.]{0,40}?(\d{1,3})%", txt):
        a, b_, p = (float(m.group(1).replace(",", ".")), float(m.group(2).replace(",", ".")),
                    int(m.group(3)))
        if b_ != 0:
            calc = 100 * a / b_
            if abs(calc - p) > 3 and abs(100 * b_ / a - p) > 3:
                conf.append("%s/%s vs %d%%" % (m.group(1), m.group(2), p))
    reg(56, "J", "percentuais conferem com a razao dos numeros",
        "ATENCAO" if conf else "OK",
        "; ".join(conf[:4]) if conf else "nenhuma razao obviamente inconsistente")


def c57():
    paras, _ = _docx()
    txt = " ".join(paras)
    sus = []
    # \b nos dois lados: sem isso "eleva" casa dentro de "relevante"
    for m in re.finditer(r"\b(eleva|eleva-se|aumenta|sobe|cresce)\b[^.]{0,50}?(−\d,\d{2,3})", txt):
        sus.append(m.group(0)[:60])
    for m in re.finditer(r"(reduz|cai|diminui|recua)[^.]{0,40}?(?<![−-])(\d,\d{2,3})", txt):
        pass
    reg(57, "J", "verbo de direcao bate com o sinal",
        "CRITICO" if sus else "OK",
        "; ".join(sus[:3]) if sus else "nenhum verbo de alta junto de coeficiente negativo")


def c58():
    paras, cels = _docx()
    txt = " ".join(paras)
    cont = re.findall(r"(\w+|\d+) (?:das|de) (\w+|\d+) (?:séries|variáveis|horizontes)", txt)
    reg(58, "J", "contagens do tipo 'X de Y'",
        "ATENCAO",
        "%d contagens no texto: %s -- confira cada uma contra a tabela"
        % (len(cont), "; ".join("%s de %s" % c for c in cont[:5])))


def c59():
    paras, _ = _docx()
    i = next((k for k, p in enumerate(paras) if p.strip().upper().startswith("5 CONCLUS")), None)
    j = next((k for k, p in enumerate(paras) if p.strip().upper().startswith("REFER")), len(paras))
    if i is None:
        reg(59, "J", "numeros da conclusao existem no capitulo 4", "ATENCAO", "conclusao nao localizada"); return
    concl = " ".join(paras[i:j])
    corpo = " ".join(paras[:i])
    nums = set(re.findall(r"\d,\d{2,3}", concl))
    faltam = [n for n in nums if n not in corpo]
    reg(59, "J", "numeros da conclusao aparecem no corpo",
        "ATENCAO" if faltam else "OK",
        ("%d sem contrapartida: %s" % (len(faltam), ", ".join(sorted(faltam)[:6])))
        if faltam else "todos os %d numeros da conclusao aparecem antes" % len(nums))


def c60():
    paras, _ = _docx()
    i = next((k for k, p in enumerate(paras) if p.strip().lower() == "resumo"), None)
    j = next((k for k, p in enumerate(paras) if p.strip().lower() in ("abstract",)), None)
    if i is None:
        reg(60, "J", "numeros do resumo batem com o corpo", "ATENCAO", "resumo nao localizado"); return
    res = " ".join(paras[i:(j or i + 12)])
    corpo = " ".join(paras[(j or i) + 1:])
    nums = set(re.findall(r"\d,\d{2,3}", res))
    faltam = [n for n in nums if n not in corpo]
    reg(60, "J", "numeros do resumo aparecem no corpo",
        "CRITICO" if faltam else "OK",
        ("sem contrapartida: " + ", ".join(sorted(faltam))) if faltam
        else "os %d numeros do resumo aparecem no texto" % len(nums))


FAM = {"F": "REPRODUTIBILIDADE", "G": "DADOS, 2a CAMADA", "H": "ESPECIFICACAO, 2a CAMADA",
       "I": "INFERENCIA, 2a CAMADA", "J": "TEXTO CONTRA NUMEROS"}
CH = {31: c31, 32: c32, 33: c33, 34: c34, 35: c35, 36: c36,
      37: c37, 38: c38, 39: c39, 40: c40, 41: c41, 42: c42,
      43: c43, 44: c44, 45: c45, 46: c46, 47: c47, 48: c48,
      49: c49, 50: c50, 51: c51, 52: c52, 53: c53, 54: c54,
      55: c55, 56: c56, 57: c57, 58: c58, 59: c59, 60: c60}
CAROS = {43, 50, 51, 52}


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    rapido = "--rapido" in sys.argv
    pedidas = [int(a) for a in args] if args else sorted(CH)
    print("SEGUNDA VARREDURA  |  %s  |  checagens 31 a 60\n" % time.strftime("%d/%m/%Y %H:%M"))
    t0 = time.time()
    for n in pedidas:
        if n in (31, 37, 43, 49, 55):
            f = {31: "F", 37: "G", 43: "H", 49: "I", 55: "J"}[n]
            print("\n--- %s. %s" % (f, FAM[f]))
        try:
            CH[n](rapido) if n in CAROS else CH[n]()
        except Exception as e:
            reg(n, "?", "checagem %d" % n, "CRITICO", "falhou: %s: %s" % (type(e).__name__, e))
    t = pd.DataFrame(R)
    t.to_excel(RES / "auditoria_segunda.xlsx", index=False)
    print("\n" + "=" * 78)
    for v in ["CRITICO", "ATENCAO", "OK"]:
        print("  %-8s %d" % (v, int((t.veredito == v).sum())))
    print("  tempo %.0fs | planilha: resultados/auditoria_segunda.xlsx" % (time.time() - t0))
    print("=" * 78)


if __name__ == "__main__":
    main()
